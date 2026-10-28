import asyncio
import io
import json

from openpyxl import Workbook

from app.services.excel_parser import parse_excel
from app.services.roster_cells import clean_roster_cell
from app.services.roster_detection import detect_roster_layouts
from app.services.student_matching import (
    StudentCandidate,
    detect_student_name,
    detect_student_name_with_llm,
    match_roster_student,
    match_student,
)


def _anonymized_roster_workbook() -> bytes:
    workbook = Workbook()
    score_sheet = workbook.active
    score_sheet.title = "Sheet1"
    score_sheet.append(
        [
            "RollNo",
            "Student Name",
            "Project",
            "Synopsis",
            "Phase1",
            "Phase2",
            "Phase3",
            "VIVA",
            "Total",
            "Remarks",
        ]
    )
    for index in range(60):
        project = f"{index + 1}. Sample Project {index + 1}" if index < 29 else ""
        score_sheet.append(
            [
                f"R-{index + 1:03d}",
                f"  SAMPLE STUDENT {index + 1:02d}  ",
                project,
                "",
                "",
                "",
                "",
                "",
                "",
                "",
            ]
        )
    score_sheet.append(["", "", "Topic Without a Student"])
    score_sheet.cell(row=1_048_576, column=1).number_format = "@"

    topic_sheet = workbook.create_sheet("Sheet2")
    topic_sheet.append(["RollNo", "Student Name", "Project TOPICS"])
    one_word_names = {
        0: "VEXARI",
        1: "QUORIUM",
        2: "ZORARI",
        3: "NEXARI",
        4: "LUMORA",
        5: "TAVIUM",
        6: "KRIORA",
        7: "YORIUM",
        8: "MOVARI",
        9: "PEXORA",
        10: "VEXORA",
    }
    for index in range(60):
        name = one_word_names.get(index, f"ANON STUDENT {index + 1:02d}")
        topic = f"\xa0\xa0Sample Topic {index + 1}"
        if index >= 12:
            topic = f"{index + 2}.  {topic}"
        topic_sheet.append(
            [
                "" if index == 25 else f"Q-{index + 1:03d}",
                f" {name} ",
                topic,
            ]
        )

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_roster_cleaner_normalizes_unicode_whitespace_and_topic_numbers():
    name = clean_roster_cell("  VEXARI\xa0\t  Quorium\n")
    numbered_topic = clean_roster_cell("\xa0Topic 5 -\tAgile Beyond IT: \xa0Future", is_topic=True)
    dotted_topic = clean_roster_cell("17.  Blockchain Attendance", is_topic=True)
    parenthesized_topic = clean_roster_cell("5) Smart Cities", is_topic=True)

    assert name.text == "VEXARI Quorium"
    assert numbered_topic.text == "Agile Beyond IT: Future"
    assert numbered_topic.topic_number == 5
    assert dotted_topic.text == "Blockchain Attendance"
    assert dotted_topic.topic_number == 17
    assert parenthesized_topic.text == "Smart Cities"
    assert parenthesized_topic.topic_number == 5


def test_anonymized_two_sheet_roster_fixture_detects_layout_and_real_header_rows():
    parsed = parse_excel(_anonymized_roster_workbook(), "anonymized-rosters.xlsx")
    layouts = detect_roster_layouts(parsed)

    assert [(layout.sheet_name, layout.header_row) for layout in layouts] == [
        ("Sheet1", 1),
        ("Sheet2", 1),
    ]
    assert all(layout.row_count == 60 for layout in layouts)
    assert all(layout.distinct_name_count == 60 for layout in layouts)
    assert layouts[0].roll_column == 0
    assert layouts[0].name_column == 1
    assert layouts[0].topic_column == 2
    assert parsed.sheets[0].real_row_count == 62
    assert 1_048_576 not in parsed.sheets[0].source_row_numbers


def test_roster_detection_reports_physical_header_row_after_blank_rows():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A3"] = "Registration No."
    worksheet["B3"] = "Student Name"
    worksheet["C3"] = "Project TOPICS"
    for index in range(3):
        worksheet.append([f"R-{index}", f"ANON NAME {index}", f"Topic {index}"])
    output = io.BytesIO()
    workbook.save(output)

    parsed = parse_excel(output.getvalue(), "header-row.xlsx")
    layouts = detect_roster_layouts(parsed)

    assert layouts[0].header_row == 3


def test_header_words_are_never_detected_as_student_names():
    header_values = (
        "Project",
        "Project TOPICS",
        "Student Name",
        "Name",
        "Roll No",
        "Topic",
        "A project topics header",
    )
    for header_value in header_values:
        workbook = Workbook()
        worksheet = workbook.active
        workbook.properties.creator = "Unknown"
        workbook.properties.lastModifiedBy = "Unknown"
        worksheet.append(["Student Name", header_value])
        output = io.BytesIO()
        workbook.save(output)
        parsed = parse_excel(output.getvalue(), "assignment.xlsx")

        assert detect_student_name(parsed, "assignment.xlsx").detected_name is None


def test_llm_header_word_is_not_accepted_as_student_name(monkeypatch):
    class FakeResponse:
        def raise_for_status(self):
            return None

        def json(self):
            return {
                "choices": [
                    {
                        "message": {
                            "content": json.dumps(
                                {
                                    "student_name": "Project TOPICS",
                                    "confidence": 1,
                                    "evidence": "header",
                                }
                            )
                        }
                    }
                ]
            }

    class FakeClient:
        def __init__(self, *, timeout):
            self.timeout = timeout

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def post(self, *_args, **_kwargs):
            return FakeResponse()

    monkeypatch.setenv("GROQ_API_KEY", "test-key")
    monkeypatch.setattr("app.services.student_matching.httpx.AsyncClient", FakeClient)
    workbook = Workbook()
    workbook.properties.creator = "Unknown"
    worksheet = workbook.active
    worksheet.append(["Instructions", "None"])
    output = io.BytesIO()
    workbook.save(output)
    parsed = parse_excel(output.getvalue(), "assignment.xlsx")

    detected = asyncio.run(
        detect_student_name_with_llm(parsed, "assignment.xlsx", use_mock=False)
    )

    assert detected.detected_name is None


def test_roster_one_word_names_require_confirmation_without_changing_upload_matching():
    candidates = [StudentCandidate(id=42, full_name="nexari")]

    roster_match = match_roster_student("nexari", candidates)
    upload_match = match_student("nexari", candidates)

    assert roster_match.match_status == "needs_confirmation"
    assert roster_match.resolved_student_id is None
    assert upload_match.match_status == "auto_matched"
    assert upload_match.resolved_student_id == 42


def _small_table(headers: list[str], names: list[str], *, signal: bool = False) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    workbook.properties.creator = "Unknown"
    if signal:
        worksheet.append(["Name:", "ANON SINGLE STUDENT"])
    worksheet.append(headers)
    for index, name in enumerate(names):
        worksheet.append([name, *(["Topic"] if len(headers) > 1 else [])])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_name_only_group_member_table_is_not_a_roster():
    parsed = parse_excel(
        _small_table(["Name"], ["ANON MEMBER A", "ANON MEMBER B", "ANON MEMBER C"]),
        "group-members.xlsx",
    )

    assert detect_roster_layouts(parsed, "group-members.xlsx") == ()


def test_name_only_attendance_list_with_fewer_than_ten_names_is_not_a_roster():
    parsed = parse_excel(
        _small_table(
            ["Name"],
            ["ANON STUDENT A", "ANON STUDENT B", "ANON STUDENT C", "ANON STUDENT D"],
        ),
        "attendance.xlsx",
    )

    assert detect_roster_layouts(parsed, "attendance.xlsx") == ()


def test_single_submission_name_signal_blocks_small_roster_classification():
    parsed = parse_excel(
        _small_table(
            ["Student Name", "Project"],
            ["ANON STUDENT A", "Sample topic"],
            signal=True,
        ),
        "single-submission.xlsx",
    )

    assert detect_roster_layouts(parsed, "single-submission.xlsx") == ()
