import io
import zipfile

import pytest

from openpyxl import Workbook

from app.services import excel_parser


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Assignment"
    worksheet.append(["Student Name", "Topic"])
    worksheet.append(["Vexari Quorium", "Blockchain"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _workbook_bytes(workbook: Workbook) -> bytes:
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_parses_xlsx_as_values_and_preserves_name_cells():
    parsed = excel_parser.parse_excel(_xlsx_bytes(), "student.xlsx")

    assert parsed.sheets[0].rows[1] == ("Vexari Quorium", "Blockchain")
    assert "Vexari Quorium" in parsed.text
    assert parsed.sheets[0].name == "Assignment"


def test_parses_csv_and_strips_control_characters():
    parsed = excel_parser.parse_excel(
        b"Student Name,Topic\nVexari\x00 Quorium,Blockchain\n",
        "student.csv",
    )

    assert parsed.sheets[0].rows[1] == ("Vexari Quorium", "Blockchain")


@pytest.mark.parametrize(
    ("filename", "content", "message"),
    [
        ("student.xlsm", b"not an excel file", "Unsupported spreadsheet type"),
        ("student.xlsx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", "Encrypted Excel"),
        ("student.xlsx", b"corrupt", "corrupt"),
    ],
)
def test_rejects_unsupported_encrypted_and_corrupt_files(filename, content, message):
    with pytest.raises(ValueError, match=message):
        excel_parser.parse_excel(content, filename)


def test_rejects_xlsx_over_uncompressed_size_limit(monkeypatch):
    monkeypatch.setattr(excel_parser, "EXCEL_MAX_UNCOMPRESSED_BYTES", 1)

    with pytest.raises(ValueError, match="uncompressed-size limit"):
        excel_parser.parse_excel(_xlsx_bytes(), "student.xlsx")


def test_rejects_zip_bomb_before_loading_workbook(monkeypatch):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", "A" * 2048)
    monkeypatch.setattr(excel_parser, "EXCEL_MAX_UNCOMPRESSED_BYTES", 128)

    with pytest.raises(ValueError, match="uncompressed-size limit"):
        excel_parser.parse_excel(output.getvalue(), "student.xlsx")


def test_workbook_over_row_limit_is_truncated(monkeypatch):
    monkeypatch.setattr(excel_parser, "EXCEL_MAX_ROWS", 1)

    parsed = excel_parser.parse_excel(_xlsx_bytes(), "student.xlsx")

    assert parsed.truncated is True
    assert parsed.real_row_count == 2
    assert parsed.retained_row_count == 1
    assert parsed.truncated_row_count == 1
    assert len(parsed.sheets[0].rows) == 1


def test_csv_over_cell_limit_is_truncated(monkeypatch):
    monkeypatch.setattr(excel_parser, "EXCEL_MAX_CELLS", 2)

    parsed = excel_parser.parse_excel(b"a,b\n1,2\n", "student.csv")

    assert parsed.truncated is True
    assert parsed.truncated_cell_count == 2
    assert parsed.sheets[0].rows == (("a", "b"), ())


def test_formatted_millionth_row_stops_after_empty_run():
    workbook = Workbook()
    worksheet = workbook.active
    for row_number in range(20):
        worksheet.append([f"Student {row_number + 1}", "Topic"])
    worksheet.cell(row=1_048_576, column=1).number_format = "@"

    parsed = excel_parser.parse_excel(_workbook_bytes(workbook), "formatted.xlsx")

    assert parsed.real_row_count == 20
    assert parsed.retained_row_count == 20
    assert len(parsed.sheets[0].rows) == 20
    assert parsed.truncated is False


def test_five_thousand_real_rows_are_truncated_with_counts(monkeypatch):
    monkeypatch.setattr(excel_parser, "EXCEL_MAX_ROWS", 2000)
    workbook = Workbook()
    worksheet = workbook.active
    for row_number in range(5000):
        worksheet.append([f"Student {row_number + 1}", "Topic"])

    parsed = excel_parser.parse_excel(_workbook_bytes(workbook), "many-rows.xlsx")

    assert parsed.truncated is True
    assert parsed.real_row_count == 5000
    assert parsed.retained_row_count == 2000
    assert parsed.truncated_row_count == 3000
    assert len(parsed.sheets[0].rows) == 2000


def test_workbook_with_too_many_sheets_is_rejected(monkeypatch):
    monkeypatch.setattr(excel_parser, "EXCEL_MAX_SHEETS", 1)
    workbook = Workbook()
    workbook.create_sheet("Second")

    with pytest.raises(excel_parser.ExcelLimitError, match="too many sheets"):
        excel_parser.parse_excel(_workbook_bytes(workbook), "many-sheets.xlsx")


def test_all_empty_sheet_is_accepted_and_stops_after_empty_run():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.cell(row=1_048_576, column=1).number_format = "@"

    parsed = excel_parser.parse_excel(_workbook_bytes(workbook), "empty.xlsx")

    assert parsed.sheets[0].rows == ()
    assert parsed.real_row_count == 0
    assert parsed.truncated is False
