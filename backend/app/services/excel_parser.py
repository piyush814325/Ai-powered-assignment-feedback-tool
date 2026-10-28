import csv
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook

from app.config import (
    EXCEL_MAX_CELLS,
    EXCEL_MAX_CONSECUTIVE_EMPTY_ROWS,
    EXCEL_MAX_ROWS,
    EXCEL_MAX_SHEETS,
    EXCEL_MAX_TEXT_CHARS,
    EXCEL_MAX_UNCOMPRESSED_BYTES,
    MAX_UPLOAD_BYTES,
)


class ExcelParseError(ValueError):
    """A safe, user-facing spreadsheet parsing error."""


class ExcelLimitError(ExcelParseError):
    """A spreadsheet exceeded a non-truncatable safety limit."""


class ExcelSafetyError(ExcelParseError):
    """A spreadsheet is unsafe or cannot be read."""


_CONTROL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_CORE_PROPERTIES_NS = {
    "dc": "http://purl.org/dc/elements/1.1/",
    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
}
_WORKBOOK_NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}


@dataclass(frozen=True)
class ParsedSheet:
    name: str
    rows: tuple[tuple[str, ...], ...]
    real_row_count: int = 0
    truncated_row_count: int = 0
    truncated_cell_count: int = 0
    source_row_numbers: tuple[int, ...] = ()


@dataclass(frozen=True)
class ParsedExcel:
    text: str
    sheets: tuple[ParsedSheet, ...]
    document_properties: dict[str, str]
    has_visual_content: bool
    truncated: bool = False
    real_row_count: int = 0
    retained_row_count: int = 0
    truncated_row_count: int = 0
    truncated_cell_count: int = 0


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    return _CONTROL_CHARACTERS.sub("", text).strip()


def _format_table_row(values: tuple[str, ...]) -> str:
    cells = (value.replace("|", "\\|").replace("\n", " ") for value in values)
    return "| " + " | ".join(cells) + " |"


def _check_sheet_limit(sheet_count: int) -> None:
    if sheet_count > EXCEL_MAX_SHEETS:
        raise ExcelLimitError(
            f"Workbook has too many sheets; the maximum is {EXCEL_MAX_SHEETS}"
        )


def _row_has_values(row: tuple[str, ...]) -> bool:
    return any(value for value in row)


def _retain_cells(row: tuple[str, ...], remaining: int) -> tuple[tuple[str, ...], int]:
    retained: list[str] = []
    cells_kept = 0
    for value in row:
        if value and cells_kept >= remaining:
            break
        retained.append(value)
        if value:
            cells_kept += 1
    return tuple(retained), cells_kept


def _parse_csv(content: bytes) -> ParsedExcel:
    try:
        text_stream = io.StringIO(content.decode("utf-8-sig"))
    except UnicodeDecodeError as error:
        raise ExcelSafetyError("CSV files must use UTF-8 encoding") from error

    rows: list[tuple[str, ...]] = []
    real_rows = 0
    retained_cells = 0
    truncated_rows = 0
    truncated_cells = 0
    empty_run = 0
    source_row_numbers: list[int] = []
    try:
        for source_row_number, row in enumerate(csv.reader(text_stream), start=1):
            cleaned = tuple(_clean_text(value) for value in row)
            if not _row_has_values(cleaned):
                empty_run += 1
                if empty_run >= EXCEL_MAX_CONSECUTIVE_EMPTY_ROWS:
                    break
                continue
            empty_run = 0
            real_rows += 1
            if len(rows) < EXCEL_MAX_ROWS:
                retained_row, kept = _retain_cells(
                    cleaned, max(0, EXCEL_MAX_CELLS - retained_cells)
                )
                retained_cells += kept
                truncated_cells += sum(bool(value) for value in cleaned) - kept
                rows.append(retained_row)
                source_row_numbers.append(source_row_number)
            else:
                truncated_rows += 1
                truncated_cells += sum(bool(value) for value in cleaned)
    except csv.Error as error:
        raise ExcelSafetyError("CSV file is malformed") from error

    sheet = ParsedSheet(
        "CSV",
        tuple(rows),
        real_row_count=real_rows,
        truncated_row_count=truncated_rows,
        truncated_cell_count=truncated_cells,
        source_row_numbers=tuple(source_row_numbers),
    )
    rendered = _render_sheets((sheet,), EXCEL_MAX_TEXT_CHARS)
    return ParsedExcel(
        rendered,
        (sheet,),
        {},
        False,
        truncated=bool(truncated_rows or truncated_cells),
        real_row_count=real_rows,
        retained_row_count=len(rows),
        truncated_row_count=truncated_rows,
        truncated_cell_count=truncated_cells,
    )


def _workbook_metadata(archive: zipfile.ZipFile) -> tuple[int, dict[str, str], bool]:
    try:
        workbook_xml = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    except (KeyError, ElementTree.ParseError) as error:
        raise ExcelSafetyError("Excel file is corrupt or is not a valid .xlsx workbook") from error
    sheet_count = len(workbook_xml.findall("main:sheets/main:sheet", _WORKBOOK_NS))
    if sheet_count == 0:
        raise ExcelSafetyError("Excel workbook contains no worksheets")
    _check_sheet_limit(sheet_count)

    properties: dict[str, str] = {}
    try:
        core_xml = ElementTree.fromstring(archive.read("docProps/core.xml"))
    except KeyError:
        core_xml = None
    except ElementTree.ParseError as error:
        raise ExcelSafetyError("Excel document properties are malformed") from error
    if core_xml is not None:
        for key in ("creator", "lastModifiedBy"):
            namespace = "dc" if key == "creator" else "cp"
            node = core_xml.find(f"{namespace}:{key}", _CORE_PROPERTIES_NS)
            if node is not None and node.text:
                properties["author" if key == "creator" else "last_modified_by"] = _clean_text(node.text)
    has_visual_content = any(
        name.startswith(("xl/drawings/", "xl/charts/", "xl/media/"))
        for name in archive.namelist()
    )
    return sheet_count, properties, has_visual_content


def _preflight_xlsx(content: bytes) -> tuple[int, dict[str, str], bool]:
    if content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        raise ExcelSafetyError("Encrypted Excel files are not supported")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ExcelSafetyError("Encrypted Excel files are not supported")
            uncompressed_size = sum(entry.file_size for entry in entries)
            if uncompressed_size > EXCEL_MAX_UNCOMPRESSED_BYTES:
                raise ExcelLimitError(
                    "Excel workbook exceeds the "
                    f"{EXCEL_MAX_UNCOMPRESSED_BYTES}-byte uncompressed-size limit"
                )
            if any(Path(entry.filename).name.lower() == "vbaproject.bin" for entry in entries):
                raise ExcelSafetyError("Macro-enabled Excel workbooks are not supported")
            return _workbook_metadata(archive)
    except zipfile.BadZipFile as error:
        raise ExcelSafetyError("Excel file is corrupt or is not a valid .xlsx workbook") from error


def _render_sheets(sheets: tuple[ParsedSheet, ...], max_chars: int) -> str:
    rendered: list[str] = []
    for sheet in sheets:
        rendered.append(f"## Sheet: {sheet.name}")
        if sheet.rows:
            rendered.append(_format_table_row(sheet.rows[0]))
            if len(sheet.rows) > 1:
                rendered.append("| " + " | ".join("---" for _ in sheet.rows[0]) + " |")
                rendered.extend(_format_table_row(row) for row in sheet.rows[1:])
        result = "\n".join(rendered)
        if len(result) >= max_chars:
            return result[:max_chars]
    return "\n".join(rendered)[:max_chars]


def parse_excel(content: bytes, filename: str) -> ParsedExcel:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".csv"}:
        raise ExcelSafetyError(
            f"Unsupported spreadsheet type: {suffix or '(no extension)'}"
        )
    if len(content) > MAX_UPLOAD_BYTES:
        raise ExcelLimitError(f"Uploaded file exceeds the {MAX_UPLOAD_BYTES}-byte limit")
    if not content:
        raise ExcelSafetyError("Spreadsheet file is empty")
    if suffix == ".csv":
        return _parse_csv(content)

    expected_sheets, properties, has_visual_content = _preflight_xlsx(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (ValueError, zipfile.BadZipFile, KeyError, OSError, EOFError) as error:
        raise ExcelSafetyError(
            "Excel file is corrupt or is not a valid .xlsx workbook"
        ) from error

    sheets: list[ParsedSheet] = []
    total_real_rows = 0
    total_retained_rows = 0
    total_truncated_rows = 0
    total_truncated_cells = 0
    retained_cells = 0
    try:
        if len(workbook.worksheets) != expected_sheets:
            raise ExcelSafetyError("Excel workbook worksheet metadata is inconsistent")
        for worksheet in workbook.worksheets:
            rows: list[tuple[str, ...]] = []
            real_rows = 0
            sheet_truncated_rows = 0
            sheet_truncated_cells = 0
            empty_run = 0
            source_row_numbers: list[int] = []
            for source_row_number, row in enumerate(
                worksheet.iter_rows(values_only=True), start=1
            ):
                cleaned = tuple(_clean_text(value) for value in row)
                if not _row_has_values(cleaned):
                    empty_run += 1
                    if empty_run >= EXCEL_MAX_CONSECUTIVE_EMPTY_ROWS:
                        break
                    continue
                empty_run = 0
                real_rows += 1
                total_real_rows += 1
                if real_rows <= EXCEL_MAX_ROWS:
                    retained_row, kept = _retain_cells(
                        cleaned, max(0, EXCEL_MAX_CELLS - retained_cells)
                    )
                    retained_cells += kept
                    discarded = sum(bool(value) for value in cleaned) - kept
                    sheet_truncated_cells += discarded
                    total_truncated_cells += discarded
                    rows.append(retained_row)
                    source_row_numbers.append(source_row_number)
                    total_retained_rows += 1
                else:
                    sheet_truncated_rows += 1
                    total_truncated_rows += 1
                    discarded = sum(bool(value) for value in cleaned)
                    sheet_truncated_cells += discarded
                    total_truncated_cells += discarded
            sheets.append(
                ParsedSheet(
                    _clean_text(worksheet.title),
                    tuple(rows),
                    real_row_count=real_rows,
                    truncated_row_count=sheet_truncated_rows,
                    truncated_cell_count=sheet_truncated_cells,
                    source_row_numbers=tuple(source_row_numbers),
                )
            )
    except (zipfile.BadZipFile, KeyError, OSError, EOFError, ElementTree.ParseError) as error:
        raise ExcelSafetyError("Excel file is corrupt or could not be read") from error
    finally:
        workbook.close()

    parsed_sheets = tuple(sheets)
    return ParsedExcel(
        _render_sheets(parsed_sheets, EXCEL_MAX_TEXT_CHARS),
        parsed_sheets,
        properties,
        has_visual_content,
        truncated=bool(total_truncated_rows or total_truncated_cells),
        real_row_count=total_real_rows,
        retained_row_count=total_retained_rows,
        truncated_row_count=total_truncated_rows,
        truncated_cell_count=total_truncated_cells,
    )


def validate_excel_upload(content: bytes, filename: str, content_type: str | None) -> None:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".csv"}:
        raise ExcelSafetyError(
            "Only .xlsx and .csv files are supported; .xlsm and .xls are not allowed"
        )
    if not content:
        raise ExcelSafetyError("Spreadsheet file is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ExcelLimitError(f"Uploaded file exceeds the {MAX_UPLOAD_BYTES}-byte limit")

    supplied_type = (content_type or "").split(";", 1)[0].strip().lower()
    generic_types = {"", "application/octet-stream", "binary/octet-stream"}
    if suffix == ".xlsx":
        expected_types = {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/zip",
        }
        if not content.startswith(b"PK\x03\x04"):
            if content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
                raise ExcelSafetyError("Encrypted Excel files are not supported")
            raise ExcelSafetyError("File content does not match an .xlsx workbook")
        if supplied_type not in generic_types | expected_types:
            raise ExcelSafetyError("File MIME type does not match the .xlsx extension")
        _preflight_xlsx(content)
        return

    if supplied_type not in generic_types | {
        "text/csv",
        "application/csv",
        "text/plain",
    }:
        raise ExcelSafetyError("File MIME type does not match the .csv extension")
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ExcelSafetyError("CSV files must use UTF-8 encoding") from error
    if "\x00" in decoded:
        raise ExcelSafetyError("CSV file contains binary data")
    _parse_csv(content)
