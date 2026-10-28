#!/usr/bin/env python3
"""Generate anonymized roster workbooks, optionally seeding matching dev students."""

import argparse
import os
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.engine import make_url

ROOT = Path(__file__).resolve().parents[1]
SAMPLE_DIR = ROOT / "samples"
WORD_A = ("Vex", "Quor", "Zor", "Nex", "Luma", "Tavi", "Krix", "Yori", "Mova", "Pex")
WORD_B = ("ora", "ium", "exa", "oni", "ari", "uvo")


def fake_word(index: int) -> str:
    return f"{WORD_A[index % len(WORD_A)]}{WORD_B[(index // len(WORD_A)) % len(WORD_B)]}{index // 60}"


def fake_name(index: int, one_word: bool = False) -> str:
    first = fake_word(index)
    if one_word:
        return first
    return f"{first} {fake_word(index + 17)}"


def style_far_empty_row(sheet) -> None:
    for column in "ABCDEFGHIJ":
        dimension = sheet.column_dimensions[column]
        dimension.font = Font(color="D9E2F3")
        dimension.fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    sheet.cell(row=1_000_000, column=1).number_format = "@"
    far_row = sheet.row_dimensions[1_000_000]
    far_row.font = Font(color="D9E2F3")
    far_row.fill = PatternFill(fill_type="solid", fgColor="FFFFFF")


def build_projects_roster(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Projects"
    sheet.append(["RollNo", "Student Name", "Project", "Synopsis", "Phase1", "Phase2", "Phase3", "VIVA", "Total", "Remarks"])
    for index in range(60):
        name = fake_name(index).upper()
        if index % 7 == 0:
            name += "  "
        topic = f"Project {index + 1}" if index % 2 == 0 else ""
        sheet.append([
            "" if index == 11 else f"FAKE-{index + 1001}",
            name,
            topic,
            "Anonymized sample synopsis" if topic else "",
            "", "", "", "", "", "",
        ])
    sheet.append(["", "", "Orphan topic", "Topic without a student", "", "", "", "", "", ""])
    style_far_empty_row(sheet)
    workbook.save(path)


def build_presentation_roster(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Presentations"
    sheet.append(["RollNo", "Student Name", "Project TOPICS"])
    for index in range(60):
        name = fake_name(index, one_word=index < 10)
        if index == 11:
            roll = ""
        else:
            roll = f"FAKE-{index + 1001}"
        if index == 14:
            prefix = "14. "
        elif index == 17:
            prefix = "17.  "
        else:
            prefix = ""
        topic = f"{prefix}{fake_word(index + 100)}\u00a0{fake_word(index + 101)}"
        sheet.append([roll, name, topic])
    style_far_empty_row(sheet)
    workbook.save(path)


def seed_development_accounts(password: str) -> None:
    from dotenv import load_dotenv

    load_dotenv(ROOT / "backend" / ".env")
    database_url = os.environ.get("DATABASE_URL", "").strip()
    production_url = os.environ.get("PRODUCTION_DATABASE_URL", "").strip()
    expected_host = os.environ.get("DEV_DATABASE_HOST", "").strip()
    if not database_url or not expected_host:
        raise SystemExit("Set DATABASE_URL and DEV_DATABASE_HOST before seeding.")
    try:
        actual_host = make_url(database_url).host
    except Exception as error:
        raise SystemExit("DATABASE_URL is invalid; refusing to seed.") from error
    if database_url == production_url or not actual_host or actual_host.casefold() != expected_host.casefold():
        raise SystemExit("Database host is not explicitly marked as development; refusing to seed.")

    sys.path.insert(0, str(ROOT / "backend"))
    from app.database import SessionLocal
    from app.models import User, UserRole
    from app.services.auth_service import get_password_hash

    db = SessionLocal()
    try:
        added = 0
        password_hash = get_password_hash(password)
        for index in range(60):
            email = f"sample-student-{index + 1:02d}@example.test"
            if db.query(User.id).filter(User.email == email).first():
                continue
            db.add(
                User(
                    email=email,
                    full_name=fake_name(index),
                    roll_no=f"FAKE-{index + 1001}",
                    hashed_password=password_hash,
                    role=UserRole.STUDENT,
                )
            )
            added += 1
        db.commit()
        print(f"Seeded {added} anonymized development student accounts.")
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--seed-dev", action="store_true", help="Create matching students in a development database.")
    parser.add_argument("--i-know-this-is-dev", action="store_true", help="Required confirmation before database seeding.")
    parser.add_argument("--student-password", default=os.environ.get("SAMPLE_STUDENT_PASSWORD", ""))
    args = parser.parse_args()
    if args.seed_dev and not args.i_know_this_is_dev:
        parser.error("--seed-dev requires --i-know-this-is-dev")
    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)
    build_projects_roster(SAMPLE_DIR / "sample_projects_roster.xlsx")
    build_presentation_roster(SAMPLE_DIR / "sample_presentation_roster.xlsx")
    print(f"Generated anonymized sample rosters in {SAMPLE_DIR.relative_to(ROOT)}.")
    if args.seed_dev:
        if not args.student_password:
            parser.error("Provide --student-password or SAMPLE_STUDENT_PASSWORD when seeding.")
        seed_development_accounts(args.student_password)


if __name__ == "__main__":
    main()
